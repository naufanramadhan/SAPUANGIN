"""Octave-only input bridge. MATLAB reads the workbook directly.

Run from any directory: python3 tools/export_workbook.py [--check]
Dependencies: openpyxl. No workbook formulas or cached values are changed.
"""
import argparse
import hashlib
import json
import math
from pathlib import Path

import openpyxl


def export(root, check=False):
    book = root / 'input' / 'FSAE_VD_3_KASAR .xlsx'
    target = root / 'input' / 'study_inputs.json'
    digest = hashlib.sha256(book.read_bytes()).hexdigest()
    if check:
        data = json.loads(target.read_text())
        if data['workbook_sha256'] != digest:
            raise SystemExit('Workbook changed: run export_workbook.py again.')
        print('Workbook SHA256 matches exported inputs.')
        return
    wb = openpyxl.load_workbook(book, data_only=True, read_only=True)
    baseline = []
    for row in range(4, 59):
        value = wb['02_INPUT_VARIABLES'].cell(row, 5).value
        if not isinstance(value, (int, float)) or not math.isfinite(value):
            raise ValueError(f'Baseline E{row} lacks a finite cached value. '
                             'Recalculate and save the workbook in Excel first.')
        baseline.append(value)
    config = {}
    for key, value in wb['06_SIMULATION_CONFIG'].iter_rows(
            min_row=2, max_col=2, values_only=True):
        if key is None:
            continue
        if key in config:
            raise ValueError(f'Duplicate config key: {key}')
        config[key] = value
    wb.close()
    data = dict(baseline=baseline, config=config,
                workbook_size_bytes=book.stat().st_size,
                workbook_mtime_unix=book.stat().st_mtime,
                workbook_sha256=digest)
    target.write_text(json.dumps(data, indent=2, allow_nan=False) + '\n')
    print(f'Exported baseline column E and configuration: {target}')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--check', action='store_true')
    args = parser.parse_args()
    export(Path(__file__).resolve().parents[1], args.check)
